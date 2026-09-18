"""Tests for the agent gym's comparison.

Every test here is built so it FAILS without the behaviour it names: a test that
passes both before and after a change measures nothing. The controls at the end
are the important ones — a clean bridge proves nothing on its own.
"""
import json
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

COMPARE = Path(__file__).parent.parent / "skills" / "agent-gym" / "scripts" / "compare.py"

EVAL_YAML = """
agent: test-agent
period: month
outputs:
  - id: reclass
    kind: workbook
    build: "true"
    candidate:
      file: output/{period}/candidate.xlsx
    reference:
      file: refs/{period}/reference.xlsx
    components:
      - id: lines
        kind: rows
        reference: {sheet: Reclass, header_row: 1, key: [txn, vendor], amount: amount,
                     fields: {target: to_account}}
        candidate: {sheet: Lines, header_row: 1, key: [txn, vendor], amount: amount,
                     fields: {target: proposed_account}, filter: {verdict: RECLASS}}
      - id: total
        kind: figure
        metric: exact
        reference: {sheet: Summary, label: Total, label_col: A, value_col: B}
        candidate: {sheet: Summary, label: Total, label_col: A, value_col: B}
eval_set:
  - {period: 2026-06, role: development}
"""


def write_book(path: Path, sheets: dict) -> None:
    """sheets: {name: [[row], [row]]}, written from A1."""
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for r, row in enumerate(rows, start=1):
            for c, value in enumerate(row, start=1):
                ws.cell(row=r, column=c, value=value)
    wb.save(path)


def build_agent(tmp_path: Path, *, candidate_rows=None, reference_rows=None,
                candidate_total=300.0, reference_total=300.0) -> Path:
    agent = tmp_path / "agent"
    (agent / "eval").mkdir(parents=True)
    (agent / "eval" / "eval.yaml").write_text(EVAL_YAML)

    candidate_rows = candidate_rows if candidate_rows is not None else [
        ["t1", "Acme", 100.0, "Travel", "RECLASS"],
        ["t2", "Globex", 200.0, "Meals", "RECLASS"],
        ["t3", "Initech", 55.0, "", "OK"],          # not in the population under test
    ]
    reference_rows = reference_rows if reference_rows is not None else [
        ["t1", "Acme", 100.0, "Travel"],
        ["t2", "Globex", 200.0, "Meals"],
    ]
    write_book(agent / "output" / "2026-06" / "candidate.xlsx", {
        "Lines": [["txn", "vendor", "amount", "proposed_account", "verdict"], *candidate_rows],
        "Summary": [["Total", candidate_total]],
    })
    write_book(agent / "refs" / "2026-06" / "reference.xlsx", {
        "Reclass": [["txn", "vendor", "amount", "to_account"], *reference_rows],
        "Summary": [["Total", reference_total]],
    })
    return agent


def compare(agent: Path, out: Path, period="2026-06"):
    result = subprocess.run(
        [sys.executable, str(COMPARE), str(agent), period, "--out", str(out), "--no-recalc"],
        capture_output=True, text=True,
    )
    bridge_path = out / f"{period}.bridge.json"
    bridge = json.loads(bridge_path.read_text()) if bridge_path.exists() else None
    return result, bridge


def refused(result, message):
    """A refusal is a stated reason and a clean exit — not a traceback that
    happens to be non-zero. Without this check an unrelated crash (a missing
    dependency, say) passes every negative test in this file."""
    output = result.stderr + result.stdout
    assert result.returncode != 0, output
    assert "Traceback" not in result.stderr, f"crashed rather than refused:\n{result.stderr}"
    assert message in output, output


def steps_of(bridge, component):
    return next(b for b in bridge["bridges"] if b["component"] == component)["steps"]


def test_matching_month_passes(tmp_path):
    result, bridge = compare(build_agent(tmp_path), tmp_path / "out")
    assert result.returncode == 0
    assert bridge["verdict"] == "pass"
    assert bridge["unexplained_steps"] == 0


def test_filter_excludes_rows_outside_the_population(tmp_path):
    """Without the filter the OK row reads as only-in-candidate and swamps the bridge."""
    _, bridge = compare(build_agent(tmp_path), tmp_path / "out")
    population = next(b for b in bridge["bridges"] if b["component"] == "lines")["population"]
    assert population == {"matched": 2, "only_reference": 0, "only_candidate": 0, "mismatched": 0}


def test_amount_difference_is_reported_to_the_cent(tmp_path):
    agent = build_agent(tmp_path, candidate_rows=[
        ["t1", "Acme", 112.34, "Travel", "RECLASS"],
        ["t2", "Globex", 200.0, "Meals", "RECLASS"],
    ])
    _, bridge = compare(agent, tmp_path / "out")
    assert bridge["verdict"] == "needs disposition"
    # keyed on txn+vendor, so the row matches and its amount differs: one step,
    # carrying the delta to the cent rather than either gross figure
    steps = steps_of(bridge, "lines")
    assert [(s["amount"], "amount differs" in s["note"]) for s in steps] == [(12.34, True)]


def test_field_difference_is_caught_without_any_dollar_delta(tmp_path):
    """Same row, same amount, different target account: no delta to give it away."""
    agent = build_agent(tmp_path, candidate_rows=[
        ["t1", "Acme", 100.0, "Office expenses", "RECLASS"],
        ["t2", "Globex", 200.0, "Meals", "RECLASS"],
    ])
    _, bridge = compare(agent, tmp_path / "out")
    notes = [s["note"] for s in steps_of(bridge, "lines")]
    assert any("target differs" in n and "Travel" in n and "Office expenses" in n for n in notes)


def test_missing_row_is_reported_with_its_amount(tmp_path):
    agent = build_agent(tmp_path, candidate_rows=[["t1", "Acme", 100.0, "Travel", "RECLASS"]])
    _, bridge = compare(agent, tmp_path / "out")
    steps = steps_of(bridge, "lines")
    assert [(s["amount"], "only in reference" in s["note"]) for s in steps] == [(200.0, True)]


def test_figure_residual(tmp_path):
    agent = build_agent(tmp_path, candidate_total=277.7)
    _, bridge = compare(agent, tmp_path / "out")
    figure = next(b for b in bridge["bridges"] if b["component"] == "total")
    assert figure["residual"] == -22.3


def test_duplicate_sheet_is_refused(tmp_path):
    """A rebuild that appends model sheets leaves stale ones in front; reading by
    name would report the old numbers as the new ones."""
    agent = build_agent(tmp_path)
    path = agent / "output" / "2026-06" / "candidate.xlsx"
    wb = load_workbook(path)
    wb.copy_worksheet(wb["Lines"]).title = "Lines1"
    wb.save(path)
    result, bridge = compare(agent, tmp_path / "out")
    refused(result, "appears 2 times")
    assert bridge is None


def test_unknown_output_kind_fails_loudly(tmp_path):
    agent = build_agent(tmp_path)
    spec = (agent / "eval" / "eval.yaml").read_text().replace("kind: workbook", "kind: google_sheet")
    (agent / "eval" / "eval.yaml").write_text(spec)
    result, _ = compare(agent, tmp_path / "out")
    refused(result, "cannot read")


def test_missing_column_names_what_it_wanted(tmp_path):
    agent = build_agent(tmp_path)
    path = agent / "refs" / "2026-06" / "reference.xlsx"
    wb = load_workbook(path)
    wb["Reclass"]["D1"] = "account"            # was to_account
    wb.save(path)
    result, _ = compare(agent, tmp_path / "out")
    refused(result, "to_account")


def test_undisposed_file_lists_every_open_difference(tmp_path):
    agent = build_agent(tmp_path, candidate_total=277.7)
    out = tmp_path / "out"
    _, bridge = compare(agent, out)
    text = (out / "2026-06.undisposed.yaml").read_text()
    for b in bridge["bridges"]:
        for step in b["steps"]:
            assert step["id"] in text
    assert "accepted_by:" in text


def test_a_disposition_settles_the_difference(tmp_path):
    agent = build_agent(tmp_path, candidate_total=277.7)
    _, bridge = compare(agent, tmp_path / "out1")
    step = steps_of(bridge, "total")[0]
    (agent / "eval" / "dispositions.yaml").write_text(
        "dispositions:\n"
        f"  - id: {step['id']}\n"
        f"    amount: {step['amount']}\n"
        "    reason: judgment\n"
        "    note: below materiality; the accountant booked it anyway\n"
        "    accepted_by: someone@example.com\n"
        "    accepted_at: 2026-09-18\n"
    )
    result, after = compare(agent, tmp_path / "out2")
    assert result.returncode == 0
    assert after["verdict"] == "pass"
    assert after["disposed_steps"] == 1
    settled = steps_of(after, "total")[0]
    assert settled["reason"] == "judgment"
    assert settled["accepted_by"] == "someone@example.com"


def test_a_disposition_stops_applying_when_the_amount_changes(tmp_path):
    """The control on disposition: a ruling covers THAT difference, not that
    line for ever."""
    agent = build_agent(tmp_path, candidate_total=277.7)
    _, bridge = compare(agent, tmp_path / "out1")
    step = steps_of(bridge, "total")[0]
    (agent / "eval" / "dispositions.yaml").write_text(
        "dispositions:\n"
        f"  - id: {step['id']}\n"
        "    reason: judgment\n"
        "    accepted_by: someone@example.com\n"
    )
    # the same figure now differs by a different amount
    path = agent / "output" / "2026-06" / "candidate.xlsx"
    wb = load_workbook(path)
    wb["Summary"]["B1"] = 250.0
    wb.save(path)
    result, after = compare(agent, tmp_path / "out2")
    assert result.returncode != 0
    assert after["verdict"] == "needs disposition"
    assert steps_of(after, "total")[0]["reason"] == "unexplained"


def test_a_bad_reason_is_refused(tmp_path):
    agent = build_agent(tmp_path, candidate_total=277.7)
    (agent / "eval" / "dispositions.yaml").write_text(
        "dispositions:\n  - id: abc123\n    reason: fine-i-guess\n    accepted_by: a@b.com\n"
    )
    result, _ = compare(agent, tmp_path / "out")
    refused(result, "is not a reason")


def test_a_disposition_without_a_name_is_refused(tmp_path):
    agent = build_agent(tmp_path)
    (agent / "eval" / "dispositions.yaml").write_text(
        "dispositions:\n  - id: abc123\n    reason: judgment\n"
    )
    result, _ = compare(agent, tmp_path / "out")
    refused(result, "accepted_by")


# ── eval.py: the walk-forward loop ───────────────────────────────────────────

EVAL_SCRIPT = Path(__file__).parent.parent / "skills" / "agent-gym" / "scripts" / "eval.py"

TWO_PERIODS = """
agent: test-agent
period: week
outputs:
  - id: reclass
    kind: workbook
    build: "true"
    candidate:
      file: output/{period}/candidate.xlsx
    reference:
      file: refs/{period}/reference.xlsx
    components:
      - id: total
        kind: figure
        metric: exact
        reference: {sheet: Summary, label: Total, label_col: A, value_col: B}
        candidate: {sheet: Summary, label: Total, label_col: A, value_col: B}
eval_set:
  - {period: 2026-w01, role: development}
  - {period: 2026-w02, role: development}
  - {period: 2026-w03, role: holdout}
"""


def build_two_period_agent(tmp_path: Path, first_differs=True) -> Path:
    agent = tmp_path / "agent"
    (agent / "eval").mkdir(parents=True)
    (agent / "eval" / "eval.yaml").write_text(TWO_PERIODS)
    for period, candidate_total in (("2026-w01", 90.0 if first_differs else 100.0),
                                    ("2026-w02", 200.0), ("2026-w03", 300.0)):
        write_book(agent / "output" / period / "candidate.xlsx",
                   {"Summary": [["Total", candidate_total]]})
    for period, reference_total in (("2026-w01", 100.0), ("2026-w02", 200.0), ("2026-w03", 300.0)):
        write_book(agent / "refs" / period / "reference.xlsx",
                   {"Summary": [["Total", reference_total]]})
    return agent


def run_eval(agent: Path, *flags):
    return subprocess.run([sys.executable, str(EVAL_SCRIPT), str(agent), *flags],
                          capture_output=True, text=True)


def periods_in_latest_report(agent: Path) -> list[str]:
    runs = sorted((agent / "eval" / "reports").iterdir())
    report = json.loads((runs[-1] / "report.json").read_text())
    return [p["period"] for p in report["periods"]]


def test_training_stops_at_the_first_period_needing_disposition(tmp_path):
    """Week 1 is settled before week 2 is opened: rules written against several
    periods at once are fitted to all of them, and none was ever a test."""
    agent = build_two_period_agent(tmp_path)
    result = run_eval(agent)
    assert result.returncode != 0
    assert periods_in_latest_report(agent) == ["2026-w01"]
    assert "settle it" in result.stdout


def test_all_flag_runs_every_development_period(tmp_path):
    agent = build_two_period_agent(tmp_path)
    run_eval(agent, "--all")
    assert periods_in_latest_report(agent) == ["2026-w01", "2026-w02"]


def test_a_clean_period_does_not_stop_the_run(tmp_path):
    agent = build_two_period_agent(tmp_path, first_differs=False)
    result = run_eval(agent)
    assert result.returncode == 0
    assert periods_in_latest_report(agent) == ["2026-w01", "2026-w02"]


def test_a_holdout_is_not_spent_without_asking(tmp_path):
    agent = build_two_period_agent(tmp_path, first_differs=False)
    run_eval(agent)
    assert not (agent / "eval" / "revealed.json").exists()
    assert "2026-w03" not in periods_in_latest_report(agent)


def test_revealing_a_holdout_records_it_as_spent(tmp_path):
    agent = build_two_period_agent(tmp_path, first_differs=False)
    run_eval(agent, "--reveal-holdout")
    assert "2026-w03" in periods_in_latest_report(agent)
    spent = json.loads((agent / "eval" / "revealed.json").read_text())
    assert "2026-w03" in spent


def test_every_output_is_built_not_only_the_first(tmp_path):
    """A second output with its own build command must run, or it is compared
    against whatever file happened to be there."""
    agent = build_two_period_agent(tmp_path, first_differs=False)
    spec = (agent / "eval" / "eval.yaml").read_text()
    second = spec[spec.index("  - id: reclass"):spec.index("eval_set:")]
    second = (second.replace("- id: reclass", "- id: second")
                    .replace('build: "true"', 'build: "echo built-second > built-second.txt"')
                    .replace("file: output/{period}/candidate.xlsx",
                             "file: output/{period}/candidate.xlsx   # same artefact, own command"))
    (agent / "eval" / "eval.yaml").write_text(spec.replace("eval_set:", second + "eval_set:"))
    # the first output's candidate already exists, so only --rebuild reaches the
    # second command; without the fix it is never run at all
    run_eval(agent, "--rebuild")
    assert (agent / "built-second.txt").exists()
