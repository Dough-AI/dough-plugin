#!/usr/bin/env python3
"""Compare one period's candidate against its reference, and write the bridge.

  uv run --with pyyaml --with openpyxl compare.py <agent-dir> <period> [--out DIR]

The bridge is the walk from reference to candidate: every difference with an
amount and a reason. Reasons are not guessed — a difference this script finds is
`unexplained` until a person labels it. An unexplained step fails its period.

Nothing here knows anything about a Stripe close. What to compare comes from the
agent's eval/eval.yaml.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import shutil
import subprocess
import sys
import tempfile
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

try:
    import yaml
except ModuleNotFoundError:                                   # pragma: no cover
    raise SystemExit(
        "PyYAML is missing, and the gym reads the agent's eval.yaml with it.\n"
        "Run this through uv, which brings its own:\n"
        "  uv run --with pyyaml --with openpyxl " + __file__
    )
from openpyxl import load_workbook


# ── prerequisites: fail loudly and by name ───────────────────────────────────
def require_soffice() -> str:
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        sys.exit(
            "LibreOffice (`soffice`) is not on PATH, and a candidate written by "
            "openpyxl carries formulas with no values. Without recalculation "
            "every figure reads blank, which is indistinguishable from a build "
            "that produced zeros. Install it, or pass --no-recalc if the "
            "candidate already holds cached values."
        )
    return exe


def recalculate(path: Path, tmp: Path) -> Path:
    exe = require_soffice()
    subprocess.run([exe, "--headless", "--convert-to", "xlsx", "--outdir", str(tmp), str(path)],
                   capture_output=True, check=False)
    out = tmp / path.name
    if not out.exists():
        sys.exit(f"LibreOffice produced no output for {path.name}")
    return out


# ── reading a workbook the way eval.yaml describes it ────────────────────────
def open_book(path: Path):
    if not path.exists():
        sys.exit(f"missing workbook: {path}")
    return load_workbook(path, data_only=True)


def assert_one_sheet_each(wb, names: list[str], which: str) -> None:
    """A rebuild that appends a second set of model sheets leaves the stale ones
    in front, and a lookup by name silently reads the old numbers. Refuse."""
    present = wb.sheetnames
    for name in names:
        dupes = [s for s in present if s == name or (s.startswith(name) and s[len(name):].isdigit())]
        if name not in present:
            sys.exit(f"{which}: no sheet named '{name}' (has: {', '.join(present)})")
        if len(dupes) > 1:
            sys.exit(
                f"{which}: '{name}' appears {len(dupes)} times ({', '.join(dupes)}). "
                "A rebuild appended new sheets without clearing the old ones; the "
                "numbers read by name would be the stale copy."
            )


def read_figure(wb, spec: dict, which: str) -> float | None:
    ws = wb[spec["sheet"]]
    if "cell" in spec:
        value = ws[spec["cell"]].value
    else:
        # label lookup: find the row whose label column holds `label`, read the
        # value column on that row. Survives a layout that grew or shrank rows.
        label, lcol, vcol = spec["label"], spec["label_col"], spec["value_col"]
        value = None
        for row in range(1, ws.max_row + 1):
            if str(ws[f"{lcol}{row}"].value or "").strip().lower() == label.lower():
                value = ws[f"{vcol}{row}"].value
                break
        else:
            sys.exit(f"{which}: no row labelled '{label}' in column {lcol} of '{spec['sheet']}'")
    if value is None:
        return None
    return coerce_amount(value, f"{which}: '{spec['sheet']}' "
                                f"{spec.get('cell') or spec.get('label')}")


def coerce_amount(raw, where: str) -> float | None:
    """A blank cell is absent; anything else must be a number we can compare."""
    if raw is None or (isinstance(raw, str) and not raw.strip()):
        return None
    if isinstance(raw, bool) or not isinstance(raw, (int, float)):
        sys.exit(
            f"{where} holds {raw!r}, which is not a number. A column formatted as "
            "Text compares as nothing at all, so this refuses rather than passing "
            "the period. Fix the cell, or drop `amount` from that component."
        )
    return float(raw)


def read_rows(wb, spec: dict, which: str) -> dict[tuple, dict]:
    """Rows keyed by the declared key columns plus an occurrence index.

    Duplicates are ordinary — two identical charges on one day are not an error —
    so the key carries which occurrence it is. Without that, two swapped rows
    read as a clean match.
    """
    ws = wb[spec["sheet"]]
    header_row = spec["header_row"]
    headers = {}
    for col in range(1, ws.max_column + 1):
        name = ws.cell(row=header_row, column=col).value
        if name is not None:
            headers[str(name).strip()] = col
    wanted = list(spec["key"]) + list(spec.get("fields", {}).values()) + list(spec.get("filter", {}))
    for needed in wanted + ([spec["amount"]] if spec.get("amount") else []):
        if needed not in headers:
            sys.exit(f"{which}: '{spec['sheet']}' has no column '{needed}' "
                     f"(row {header_row} holds: {', '.join(sorted(headers))})")
    # A candidate often carries more than the population under test — every line
    # it looked at, not only the ones it acted on. `filter` says which rows count;
    # without it the rest read as "only in candidate" and swamp the bridge.
    row_filter = {headers[col]: value for col, value in spec.get("filter", {}).items()}
    fields = {label: headers[col] for label, col in spec.get("fields", {}).items()}

    out: dict[tuple, dict] = {}
    seen: Counter = Counter()
    for row in range(header_row + 1, ws.max_row + 1):
        key_parts = tuple(ws.cell(row=row, column=headers[k]).value for k in spec["key"])
        if all(p is None for p in key_parts):
            continue
        if any(str(ws.cell(row=row, column=col).value) != str(value)
               for col, value in row_filter.items()):
            continue
        seen[key_parts] += 1
        key = key_parts + (seen[key_parts],)
        amount = None
        if spec.get("amount"):
            raw = ws.cell(row=row, column=headers[spec["amount"]]).value
            # A text amount used to become None, and a None amount skips the
            # comparison — so a column formatted as Text silently disabled it and
            # the period passed. Refuse instead, naming the cell.
            amount = coerce_amount(raw, f"{which}: '{spec['sheet']}' row {row}, "
                                        f"column '{spec['amount']}'")
        out[key] = {"row": row, "amount": amount,
                    "fields": {label: ws.cell(row=row, column=col).value
                               for label, col in fields.items()}}
    return out


# ── readers: the only code that knows a file format ──────────────────────────
#
# A reader turns one artefact into the two things a bridge needs — named figures
# and keyed rows — so `compare` itself stays format-blind. Adding a Google Sheet,
# a JSON summary or a lake table means adding a reader, not touching the bridge.
def read_workbook(path: Path, components: list[dict], side: str, tmp: Path,
                  recalc: bool) -> dict:
    """An .xlsx, read through openpyxl. Recalculated first when asked, because a
    workbook written by openpyxl carries formulas with no values."""
    live = recalculate(path, tmp) if recalc else path
    wb = open_book(live)
    assert_one_sheet_each(wb, [c[side]["sheet"] for c in components], side)
    figures, rows = {}, {}
    for component in components:
        spec = component[side]
        if component["kind"] == "figure":
            figures[component["id"]] = read_figure(wb, spec, side)
        elif component["kind"] == "rows":
            rows[component["id"]] = read_rows(wb, spec, side)
        else:
            sys.exit(f"unknown component kind: {component['kind']}")
    return {"figures": figures, "rows": rows}


READERS = {"workbook": read_workbook}


def merge_override(component: dict, overrides: dict) -> dict:
    patch = overrides.get(component["id"])
    if not patch:
        return component
    return {**component, **{k: {**component.get(k, {}), **v} if isinstance(v, dict) else v
                            for k, v in patch.items()}}


def build_bridge(component: dict, ref: dict, cand: dict) -> dict:
    """Dispatch on the component's kind. The bridge logic below is shared by
    every output kind: only the reader differs."""
    cid = component["id"]
    if component["kind"] == "figure":
        return figure_bridge(component, ref["figures"][cid], cand["figures"][cid])
    return rows_bridge(component, ref["rows"][cid], cand["rows"][cid])


# ── dispositions: a person's ruling on a difference, carried across runs ─────
#
# A step's id is what a disposition names. It includes the AMOUNT on purpose: if
# the same line later differs by a different amount, that is a new difference and
# must resurface rather than inherit a ruling made about the old one.
DISPOSITION_REASONS = ("judgment", "reference_error", "stale_data", "bug")

# `bug` means the agent is wrong. Ruling it does not settle the period: the fix
# does, and the next run is the evidence.
SETTLING_REASONS = ("judgment", "reference_error", "stale_data")

# How many steps of each kind a bridge lists. The counts in `population` are
# always complete; `truncated_steps` says how many were left out, and a bridge
# with any is never a pass.
STEP_CAP = 20


def step_id(period: str, output: str, component: str, detail: str, amount) -> str:
    """`detail` must identify the step WITHIN its component.

    It used to be the human-readable note, which for a row step carried only the
    first key column — so two different rows sharing it produced one id, and a
    single ruling settled both. Callers now pass the full key, its occurrence
    index and the field.
    """
    raw = "|".join([period, output, component, detail,
                    "" if amount is None else f"{float(amount):.2f}"])
    return hashlib.sha256(raw.encode()).hexdigest()[:16]


def load_dispositions(agent: Path) -> dict[str, dict]:
    path = agent / "eval" / "dispositions.yaml"
    if not path.exists():
        return {}
    loaded = yaml.safe_load(path.read_text()) or {}
    out = {}
    for entry in loaded.get("dispositions", []):
        missing = [f for f in ("id", "reason", "accepted_by") if not entry.get(f)]
        if missing:
            sys.exit(f"dispositions.yaml: an entry is missing {', '.join(missing)}")
        if entry["reason"] not in DISPOSITION_REASONS:
            sys.exit(f"dispositions.yaml: '{entry['reason']}' is not a reason "
                     f"({', '.join(DISPOSITION_REASONS)})")
        out[entry["id"]] = entry
    return out


def apply_dispositions(bridges: list[dict], period: str, rulings: dict) -> None:
    """Label each step that has a ruling. Everything else stays unexplained, and
    an unexplained step is what fails a period."""
    for bridge in bridges:
        for step in bridge["steps"]:
            step["id"] = step_id(period, bridge.get("output", ""), bridge["component"],
                                 step.get("detail", step["note"]), step["amount"])
            ruling = rulings.get(step["id"])
            if ruling:
                step["reason"] = ruling["reason"]
                step["settles"] = ruling["reason"] in SETTLING_REASONS
                step["accepted_by"] = ruling["accepted_by"]
                accepted_at = ruling.get("accepted_at")
                step["accepted_at"] = accepted_at.isoformat() if hasattr(accepted_at, "isoformat") else accepted_at
                if ruling.get("note"):
                    step["disposition_note"] = ruling["note"]
        # A bridge passes only when nothing is left unexplained AND nothing was
        # hidden by the step cap. Judging by the visible steps alone let 60
        # missing rows read as a pass once the first 20 were disposed.
        undisposed = [s for s in bridge["steps"]
                      if s["reason"] == "unexplained" or s.get("settles") is False]
        bridge["passed"] = not undisposed and not bridge.get("truncated_steps")


def write_undisposed(out_dir: Path, period: str, bridges: list[dict]) -> Path | None:
    """The review surface: every step still unexplained, ready for a person to
    add a reason and their name, then paste into eval/dispositions.yaml."""
    open_steps = [(b, s) for b in bridges for s in b["steps"] if s["reason"] == "unexplained"]
    if not open_steps:
        return None
    lines = [
        f"# {period}: {len(open_steps)} difference(s) still unexplained.",
        "#",
        "# Give each one a reason and your name, then move it into",
        "# eval/dispositions.yaml. Reasons: " + ", ".join(DISPOSITION_REASONS) + ".",
        "#   judgment        the person decided differently, and that is allowed",
        "#   reference_error the human's workbook is wrong",
        "#   stale_data      the sources moved after the reference was built",
        "#   bug             the agent is wrong — fix it rather than accept it",
        "",
        "dispositions:",
    ]
    for bridge, step in open_steps:
        lines += [
            f"  # {bridge['component']}: {step['note']}",
            f"  - id: {step['id']}",
            f"    amount: {'null' if step['amount'] is None else step['amount']}",
            "    reason:            # one of " + " | ".join(DISPOSITION_REASONS),
            "    note:",
            "    accepted_by:",
            "    accepted_at:",
            "",
        ]
    path = out_dir / f"{period}.undisposed.yaml"
    path.write_text("\n".join(lines))
    return path


# ── the bridge ───────────────────────────────────────────────────────────────
def figure_bridge(component: dict, ref: float | None, cand: float | None) -> dict:
    metric = component.get("metric", "exact")
    tolerance = float(component.get("amount", 0)) if metric == "within" else 0.0
    steps = []
    if ref is None or cand is None:
        steps.append({"amount": None, "reason": "unexplained", "detail": "empty_side",
                      "note": f"reference={ref} candidate={cand} — one side is empty"})
        residual = None
    else:
        delta = round(cand - ref, 2)
        if abs(delta) > tolerance:
            steps.append({"amount": delta, "reason": "unexplained", "detail": "residual",
                          "note": "candidate minus reference"})
        residual = delta
    return {
        "component": component["id"], "kind": "figure",
        "reference": ref, "candidate": cand,
        "steps": steps, "residual": residual,
        "passed": bool(steps == []),
    }


def rows_bridge(component: dict, ref: dict, cand: dict) -> dict:
    only_ref = sorted(set(ref) - set(cand))
    only_cand = sorted(set(cand) - set(ref))
    mismatched = []
    for key in set(ref) & set(cand):
        a, b = ref[key]["amount"], cand[key]["amount"]
        if a is not None and b is not None and round(b - a, 2) != 0:
            mismatched.append({"key": [str(p) for p in key], "field": "amount",
                               "reference": a, "candidate": b, "difference": round(b - a, 2)})
        # Declared fields: a row can be in both populations and still disagree —
        # the same charge reclassified to a different account is a real
        # difference, and it carries no dollar delta to give it away.
        for label, ref_value in (ref[key].get("fields") or {}).items():
            cand_value = (cand[key].get("fields") or {}).get(label)
            if str(ref_value or "").strip() != str(cand_value or "").strip():
                mismatched.append({"key": [str(p) for p in key], "field": label,
                                   "reference": ref_value, "candidate": cand_value,
                                   "difference": None})
    def detail_of(key, field="row"):
        return "|".join(["" if p is None else str(p) for p in key] + [field])

    steps = []
    for key in only_ref[:STEP_CAP]:
        steps.append({"amount": ref[key]["amount"], "reason": "unexplained",
                      "detail": "only_reference|" + detail_of(key),
                      "note": f"only in reference: {' '.join(str(p) for p in key[:-1])}"})
    for key in only_cand[:STEP_CAP]:
        steps.append({"amount": cand[key]["amount"], "reason": "unexplained",
                      "detail": "only_candidate|" + detail_of(key),
                      "note": f"only in candidate: {' '.join(str(p) for p in key[:-1])}"})
    for m in mismatched[:STEP_CAP]:
        who = " ".join(m["key"][:-1])
        if m["field"] == "amount":
            note = f"amount differs on {who}"
        else:
            note = (f"{m['field']} differs on {who}: "
                    f"reference {m['reference']!r} vs candidate {m['candidate']!r}")
        steps.append({"amount": m["difference"], "reason": "unexplained",
                      "detail": "mismatch|" + "|".join(m["key"]) + "|" + m["field"],
                      "note": note})
    return {
        "component": component["id"], "kind": "rows",
        "population": {"matched": len(set(ref) & set(cand)), "only_reference": len(only_ref),
                       "only_candidate": len(only_cand), "mismatched": len(mismatched)},
        "steps": steps,
        "truncated_steps": max(0, len(only_ref) + len(only_cand) + len(mismatched) - len(steps)),
        "passed": not (only_ref or only_cand or mismatched),
    }


def next_run_dir(reports: Path) -> Path:
    """A run counter, not a hash: the same agent can be evaluated twice — the
    lake resynced, or agent mode answered differently — and neither report may
    overwrite the other."""
    reports.mkdir(parents=True, exist_ok=True)
    used = [int(p.name) for p in reports.iterdir() if p.is_dir() and p.name.isdigit()]
    run = max(used, default=0) + 1
    path = reports / f"{run:04d}"
    path.mkdir()
    return path


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("agent_dir")
    ap.add_argument("period")
    ap.add_argument("--out", help="report directory (default: <agent>/eval/reports/<run>)")
    ap.add_argument("--no-recalc", action="store_true")
    args = ap.parse_args()

    agent = Path(args.agent_dir).expanduser().resolve()
    spec = yaml.safe_load((agent / "eval" / "eval.yaml").read_text())
    period = args.period

    case = next((c for c in spec["eval_set"] if c["period"] == period), None)
    if case is None:
        sys.exit(f"{period} is not in this agent's eval set")

    outputs = spec["outputs"]
    if not outputs:
        sys.exit("eval.yaml declares no outputs")

    # A reference is a human artefact and its layout moves: Atticus's own Summary
    # Check put the holding balance on row 26 in January and row 27 from February
    # on. An override says so explicitly, per period, rather than letting a fixed
    # cell read blank and call it zero.
    overrides = case.get("overrides", {})

    bridges: list[dict] = []
    compared: list[dict] = []
    with tempfile.TemporaryDirectory() as tmpdir:
        tmp = Path(tmpdir)
        for index, output in enumerate(outputs):
            kind = output.get("kind", "workbook")
            reader = READERS.get(kind)
            if reader is None:
                sys.exit(
                    f"output '{output.get('id', index)}' is kind '{kind}', which this "
                    f"gym cannot read. Known kinds: {', '.join(sorted(READERS))}."
                )

            cand_path = agent / output["candidate"]["file"].format(period=period)
            # A reference may sit outside the workspace (the usual case — it keeps
            # it unreadable during a build) or inside it. Resolve a relative path
            # against the agent, never against whatever directory the gym ran from.
            ref_path = Path(output["reference"]["file"].format(period=period)).expanduser()
            if not ref_path.is_absolute():
                ref_path = agent / ref_path
            components = [merge_override(c, overrides) for c in output["components"]]

            cand = reader(cand_path, components, "candidate", tmp,
                          recalc=not args.no_recalc and bool(output["candidate"].get("recalculate")))
            ref = reader(ref_path, components, "reference", tmp, recalc=False)

            for component in components:
                bridge = build_bridge(component, ref, cand)
                bridge["output"] = output.get("id", str(index))
                bridges.append(bridge)
            compared.append({"output": output.get("id", str(index)), "kind": kind,
                             "candidate": str(cand_path), "reference": str(ref_path)})

    rulings = load_dispositions(agent)
    apply_dispositions(bridges, period, rulings)
    # What fails a period: a difference nobody has ruled on, a difference ruled a
    # `bug` (the fix settles it, not the ruling), and differences the step cap
    # hid — the verdict has to see the same things `bridge["passed"]` does.
    unexplained = sum(1 for b in bridges for s in b["steps"] if s["reason"] == "unexplained")
    open_bugs = sum(1 for b in bridges for s in b["steps"] if s.get("settles") is False)
    hidden = sum(b.get("truncated_steps", 0) for b in bridges)
    disposed = sum(1 for b in bridges for s in b["steps"]
                   if s["reason"] != "unexplained" and s.get("settles") is not False)
    settled = unexplained == 0 and open_bugs == 0 and hidden == 0
    report = {
        "agent": spec.get("agent", agent.name),
        "period": period,
        "role": case.get("role"),
        "outputs": compared,
        "compared_at": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "verdict": "pass" if settled else "needs disposition",
        "unexplained_steps": unexplained,
        "disposed_steps": disposed,
        "open_bugs": open_bugs,
        "hidden_differences": hidden,
        "bridges": bridges,
    }

    out_dir = Path(args.out) if args.out else next_run_dir(agent / "eval" / "reports")
    out_dir.mkdir(parents=True, exist_ok=True)
    (out_dir / f"{period}.bridge.json").write_text(json.dumps(report, indent=2) + "\n")

    tally = [f"{unexplained} unexplained"]
    if open_bugs:
        tally.append(f"{open_bugs} bug(s) to fix")
    if hidden:
        tally.append(f"{hidden} more not listed")
    print(f"{period} — {report['verdict']} ({', '.join(tally)})")
    multi = len(compared) > 1
    for b in bridges:
        name = f"{b['output']}/{b['component']}" if multi else b["component"]
        if b["kind"] == "figure":
            print(f"  {name:<22} reference {b['reference']!s:>16}  candidate {b['candidate']!s:>16}"
                  f"  residual {b['residual']!s:>10}")
        else:
            p = b["population"]
            print(f"  {name:<22} matched {p['matched']:>6}  only-ref {p['only_reference']:>4}"
                  f"  only-cand {p['only_candidate']:>4}  value-diffs {p['mismatched']:>4}")
    open_file = write_undisposed(out_dir, period, bridges)
    print(f"\nbridge written to {out_dir / (period + '.bridge.json')}")
    if disposed:
        print(f"{disposed} difference(s) already disposed")
    if open_bugs:
        print(f"{open_bugs} ruled a bug — fix the agent and rerun; a ruling does not settle one")
    if hidden:
        print(f"{hidden} difference(s) beyond the first {STEP_CAP} per kind are not listed "
              "individually; the population counts are complete")
    if open_file:
        print(f"{unexplained} still unexplained — fill in and merge: {open_file}")
    return 0 if settled else 1


if __name__ == "__main__":
    raise SystemExit(main())
